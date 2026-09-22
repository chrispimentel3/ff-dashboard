"""Expected wins — the power-ranking model from Chris's "Expected Wins" workbook.

A fantasy record is mostly schedule. Score the third-highest total in the league and lose
because you drew the highest; score badly and win because you drew worse. Expected wins
throws the schedule away and asks what a team's scores were worth against the *field*.

For each week:

    scale        = sample stdev of that week's scores ÷ 1.7
    P(i beats j) = 1 / (1 + exp(-(score_i - score_j) / scale))
    xWins_i      = mean of that probability over the other 11 teams

The logistic is the point. An all-play record counts a half-point win the same as a
sixty-point win; this says a half-point win is a coin flip and a blowout is nearly certain.
The 1.7 makes the logistic approximate a normal curve, so "how far apart were these two
scores" is read against how spread out scoring actually was that week — a 20-point gap in a
tight week means more than the same gap in a wild one.

Season totals: xWins sums the weekly figures, and luck is actual wins minus xWins. A team
two wins above its xWins has been carried by the schedule and should expect a slide.

    python -m mega.xwins        # check the implementation against the workbook's own numbers
"""
from __future__ import annotations

import math
import statistics

import pandas as pd

DIVISOR = 1.7          # makes the logistic approximate a normal CDF
MIN_TEAMS = 3          # below this a week's spread means nothing


def week_xwins(scores: dict[str, float]) -> dict[str, float]:
    """One week of scores -> each team's expected wins that week (0..1)."""
    teams = [t for t, v in scores.items() if v is not None and not pd.isna(v)]
    if len(teams) < MIN_TEAMS:
        return {}
    vals = [float(scores[t]) for t in teams]
    scale = statistics.stdev(vals) / DIVISOR       # sample stdev, as STDEV.S
    if scale <= 0:
        return {t: 0.5 for t in teams}             # everyone tied: every game a coin flip
    out = {}
    for t in teams:
        s = float(scores[t])
        p = [1 / (1 + math.exp(-((s - float(scores[o])) / scale))) for o in teams if o != t]
        out[t] = sum(p) / len(p)
    return out


def week_scale(scores: dict[str, float]) -> float:
    vals = [float(v) for v in scores.values() if v is not None and not pd.isna(v)]
    return statistics.stdev(vals) / DIVISOR if len(vals) >= MIN_TEAMS else float("nan")


def weekly_table(scores: pd.DataFrame, team_col: str = "team", week_col: str = "week",
                 points_col: str = "points") -> pd.DataFrame:
    """Long frame of (team, week, points) -> the same rows plus that week's xwins and rank."""
    rows = []
    for wk, g in scores.dropna(subset=[points_col]).groupby(week_col):
        by_team = dict(zip(g[team_col], g[points_col]))
        xw = week_xwins(by_team)
        if not xw:
            continue
        order = sorted(by_team, key=lambda t: -float(by_team[t]))
        rank = {t: i + 1 for i, t in enumerate(order)}
        for t, pts in by_team.items():
            rows.append({team_col: t, week_col: wk, points_col: float(pts),
                         "xwins": xw.get(t), "rank": rank[t], "scale": week_scale(by_team)})
    return pd.DataFrame(rows)


def power_table(scores: pd.DataFrame, standings: pd.DataFrame | None = None,
                team_col: str = "team", week_col: str = "week",
                points_col: str = "points") -> pd.DataFrame:
    """Season power rankings: xWins, expected win %, luck against the real record, consistency.

    `power` is xWins per week **played**, not per scheduled week — the workbook divides by the
    full 14 either way, which reads as though everyone is terrible in September. Same number
    once a season finishes; honest before that.
    """
    wk = weekly_table(scores, team_col, week_col, points_col)
    if wk.empty:
        return pd.DataFrame()
    t = wk.groupby(team_col).agg(
        games=(week_col, "nunique"), xwins=("xwins", "sum"),
        pf=(points_col, "sum"), ppg=(points_col, "mean"),
        best=(points_col, "max"), worst=(points_col, "min"),
        sd=(points_col, lambda s: statistics.stdev(s) if len(s) > 1 else 0.0),
    ).reset_index()
    t["power"] = t["xwins"] / t["games"].replace(0, pd.NA)
    t["cv"] = t["sd"] / t["ppg"].replace(0, pd.NA)          # lower = steadier week to week

    # Actual record, preferably from the same rows: a standings file scraped a week earlier
    # would be compared against xWins over a different number of games, and every team would
    # look unlucky.
    if "opp_points" in scores.columns:
        g = scores.dropna(subset=[points_col, "opp_points"])
        rec = g.assign(w=(g[points_col] > g["opp_points"]).astype(float)
                       + 0.5 * (g[points_col] == g["opp_points"]))
        t = t.merge(rec.groupby(team_col)["w"].sum().rename("wins").reset_index(), on=team_col, how="left")
    elif standings is not None and not standings.empty and "wins" in standings.columns:
        keep = [c for c in (team_col, "wins", "losses", "ties") if c in standings.columns]
        t = t.merge(standings[keep], on=team_col, how="left")
    if "wins" in t.columns:
        t["luck"] = t["wins"] - t["xwins"]                  # + = the schedule has been kind
    t["power_rank"] = t["power"].rank(ascending=False, method="min")
    return t.sort_values("power", ascending=False).reset_index(drop=True)


if __name__ == "__main__":
    import sys
    from pathlib import Path

    book = Path(sys.argv[1] if len(sys.argv) > 1
                else "~/Downloads/2025 Expected Wins - v2.xlsx").expanduser()
    import openpyxl

    wb = openpyxl.load_workbook(book, data_only=True)
    ws = wb["Weekly Scores"]
    head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=13):
        team = r[0].value
        if not team:
            continue
        for i, c in enumerate(r[1:15], start=1):
            if isinstance(c.value, (int, float)):
                rows.append({"team": team, "week": i, "points": float(c.value)})
    scores = pd.DataFrame(rows)
    print(f"{book.name}: {scores['team'].nunique()} teams, weeks {sorted(scores['week'].unique())}")

    # The workbook's own numbers are the oracle.
    want_week1 = {c[0].value: c[17].value for c in wb["Week 1"].iter_rows(min_row=3, max_row=14)}
    got_week1 = week_xwins(dict(zip(scores[scores.week == 1]["team"], scores[scores.week == 1]["points"])))
    worst = max(abs(got_week1[t] - want_week1[t]) for t in want_week1)
    print(f"week 1 scale: {week_scale(dict(zip(scores[scores.week==1]['team'], scores[scores.week==1]['points']))):.4f}"
          f"  (workbook 12.1335)")
    print(f"week 1 xwins: largest difference from the workbook = {worst:.2e}")

    summary = {r[0].value: (r[2].value, r[4].value) for r in wb["Summary"].iter_rows(min_row=2, max_row=13)}
    got = power_table(scores).set_index("team")
    bad = []
    for team, (xw, cv) in summary.items():
        d_xw = abs(got.at[team, "xwins"] - xw)
        d_cv = abs(round(got.at[team, "cv"], 2) - round(cv, 2))
        if d_xw > 1e-6 or d_cv > 0.011:
            bad.append(f"{team}: xwins {got.at[team,'xwins']:.4f} vs {xw:.4f}, cv {got.at[team,'cv']:.3f} vs {cv:.3f}")
    print("season xWins + CV match the workbook" if not bad else "MISMATCH:\n  " + "\n  ".join(bad))
    print()
    print(got.reset_index()[["team", "games", "xwins", "power", "pf", "cv"]]
          .to_string(index=False, float_format=lambda v: f"{v:.3f}"))
